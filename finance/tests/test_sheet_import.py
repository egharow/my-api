from datetime import date, datetime

import openpyxl
import pytest

from ftracker import accounts, balances, expected, fx, sheet_import, summary
from .helpers import bank_file, run_import


def _put(ws, row, col, value, fmt=None):
    c = ws.cell(row=row + 1, column=col + 1, value=value)
    if fmt:
        c.number_format = fmt


def _month_tab(wb, title, rows, income=(("Ely", 12500, 20000), ("Shir", 7900, 8000)), mortgage=6741):
    ws = wb.create_sheet(title)
    _put(ws, 17, 1, "INCOME"); _put(ws, 17, 8, "DEBT"); _put(ws, 17, 16, "BILLS"); _put(ws, 17, 24, "EXPENSES")
    for i, (name, exp, act) in enumerate(income):
        _put(ws, 19 + i, 1, name); _put(ws, 19 + i, 4, exp); _put(ws, 19 + i, 6, act)
    _put(ws, 21, 1, "TOTAL")
    _put(ws, 31, 1, "SAVINGS")
    _put(ws, 33, 1, "Fund A"); _put(ws, 33, 4, 700); _put(ws, 33, 6, 1000)
    _put(ws, 44, 1, "TOTAL")
    _put(ws, 19, 9, "Mortgage"); _put(ws, 19, 12, 6300); _put(ws, 19, 14, mortgage)
    _put(ws, 44, 8, "TOTAL")
    _put(ws, 19, 17, "Phone"); _put(ws, 19, 20, 22); _put(ws, 19, 22, 29.89)
    _put(ws, 44, 16, "TOTAL")
    _put(ws, 19, 24, "Food shopping"); _put(ws, 19, 26, 1500); _put(ws, 19, 28, 999)
    _put(ws, 32, 31, "Bills")
    _put(ws, 46, 8, "EXPENSES LOG"); _put(ws, 47, 8, "DATE")
    for i, (d, amount, desc, cat) in enumerate(rows):
        _put(ws, 48 + i, 8, d); _put(ws, 48 + i, 10, "₪"); _put(ws, 48 + i, 11, amount)
        _put(ws, 48 + i, 13, desc); _put(ws, 48 + i, 20, cat)
    return ws


def _workbook(path, with_net_worth=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _month_tab(wb, "Tab called February", [
        (datetime(2025, 3, 3), 100.0, "SUPER EXAMPLE 12", "Food shopping"),
        (datetime(2025, 3, 9), 80.0, "SUPER EXAMPLE 44", "Food shopping"),
        (datetime(2025, 3, 12), 30.0, "ODD SHOP", "Fun"),
        (datetime(2025, 3, 13), 25.0, "ODD SHOP", "Gifts"),
        (datetime(2025, 3, 14), 300.0, "בנהפ BIT העברה ב", "Eating out"),
        (datetime(2025, 3, 15), -18.9, "STORE REFUND", "Food shopping"),
        (datetime(2025, 3, 16), 50.0, "MYSTERY", None),
        (datetime(2025, 2, 27), 20.0, "STRAGGLER", "Fun"),
        (None, 2900.0, "NO DATE ITEM", "Education"),
        (datetime(2025, 3, 20), 7.0, "NEW CATEGORY", "Vituri"),
    ])
    _month_tab(wb, "Mar again", [(datetime(2025, 3, 5), 10.0, "DUPLICATE MONTH", "Fun")])
    ws = wb.create_sheet("Net worth")
    _put(ws, 0, 4, datetime(2024, 10, 5)); _put(ws, 0, 5, "17/09/25")        # 10 May 2024 and 17 Sep 2025
    _put(ws, 1, 2, "בנקים")
    _put(ws, 2, 3, "לאומי"); _put(ws, 2, 4, 31727, "[$ ₪]#,##0"); _put(ws, 2, 5, 39171, "[$ ₪]#,##0")
    _put(ws, 3, 2, "מניות חול")
    _put(ws, 4, 3, "ETRADE"); _put(ws, 4, 4, 92196, '"$"#,##0'); _put(ws, 4, 5, 133000, '"$"#,##0')
    _put(ws, 5, 2, "שיר וילדים פנסיוני")
    _put(ws, 6, 3, "שיר פנסיה כלל"); _put(ws, 6, 5, 51114, "[$ ₪]#,##0")
    _put(ws, 7, 3, "שיר קרן השתלמות"); _put(ws, 7, 5, 20251, "General")
    hl = wb.create_sheet("High level")
    _put(hl, 1, 1, "ינואר"); _put(hl, 1, 2, "פברואר")
    _put(hl, 2, 0, "הלוואות"); _put(hl, 2, 2, 867); _put(hl, 3, 2, 1142)
    _put(hl, 5, 0, "משכנתא"); _put(hl, 5, 1, 6686)
    wb.create_sheet("Expenses")
    wb.save(path)
    return path


@pytest.fixture
def sheet(tmp_path):
    return _workbook(tmp_path / "history.xlsx")


def test_parse_reads_month_from_dates_not_tab_names_and_skips_duplicate_months(sheet):
    data = sheet_import.parse_workbook(sheet)
    assert [(m.month, m.tab) for m in data.months] == [("2025-03", "Tab called February")]
    assert any("both 2025-03" in w for w in data.warnings)
    assert any("no date" in w for w in data.warnings)
    assert any("Expenses" in w for w in data.warnings)
    m = data.months[0]
    assert next(l for l in m.log if l.description == "NO DATE ITEM").day == date(2025, 3, 1)


def test_dashboard_blocks_become_entries_and_bills_are_not_double_counted(sheet):
    m = sheet_import.parse_workbook(sheet).months[0]
    by = {(e.section, e.label): e for e in m.entries}
    assert (by[("income", "Ely")].expected, by[("income", "Ely")].actual) == (12500, 20000)
    assert by[("saving", "Fund A")].actual == 1000
    assert by[("debt", "Mortgage")].actual == 6741
    assert by[("budget", "Phone")].expected == 22 and by[("budget", "Phone")].actual is None
    assert by[("budget", "Food shopping")].expected == 1500       # target only; actuals come from the log


def test_net_worth_dates_are_read_day_first_and_currency_comes_from_the_format(sheet):
    data = sheet_import.parse_workbook(sheet)
    assert data.net_worth_dates == [date(2024, 5, 10), date(2025, 9, 17)]
    rows = {r.name: r for r in data.net_worth}
    assert rows["ETRADE"].values[date(2025, 9, 17)] == (133000.0, "USD")
    assert rows["לאומי"].values[date(2024, 5, 10)] == (31727.0, "ILS")
    assert any("no currency" in w for w in data.warnings)


def test_preview_lists_unmapped_categories_and_writes_nothing(conn, sheet):
    data = sheet_import.parse_workbook(sheet)
    text = sheet_import.preview(data)
    assert "'Vituri'" in text and "2025-03" in text
    assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 0


def test_save_keeps_your_categories_applies_your_rules_and_is_repeatable(conn, sheet):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    data = sheet_import.parse_workbook(sheet)
    res = sheet_import.save(conn, data)
    assert res["transactions"] == 10 and res["transfers_reclassified"] == 1
    rows = {r["description"]: r for r in conn.execute(
        "SELECT t.*, c.name AS cat FROM transactions t JOIN categories c ON c.id = t.category_id")}
    assert (rows["SUPER EXAMPLE 12"]["cat"], rows["SUPER EXAMPLE 12"]["category_status"]) == ("Groceries", "approved")
    assert rows["SUPER EXAMPLE 12"]["amount"] == -100.0 and rows["SUPER EXAMPLE 12"]["budget_month"] == "2025-03"
    assert rows["STORE REFUND"]["amount"] == 18.9 and rows["STORE REFUND"]["kind"] == "refund"
    bit = rows["בנהפ BIT העברה ב"]
    assert (bit["cat"], bit["kind"]) == ("Transfer to household", "transfer")      # not "Eating out"
    assert rows["MYSTERY"]["category_status"] == "proposed" and rows["MYSTERY"]["cat"] == "Uncategorised"
    assert rows["NEW CATEGORY"]["cat"] == "Vituri"                                 # unmapped kept, not lost
    again = sheet_import.save(conn, data)
    assert again["transactions"] == 0 and again["batch_id"] is None
    assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 10
    assert conn.execute("SELECT COUNT(*) FROM batches").fetchone()[0] == 1       # no empty draft left behind


def test_history_teaches_rules_for_new_statements_but_not_for_ambiguous_or_bit(conn, sheet, home, monkeypatch):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    sheet_import.save(conn, sheet_import.parse_workbook(sheet))
    patterns = {r["pattern"]: r["name"] for r in conn.execute(
        "SELECT r.pattern, c.name FROM rules r JOIN categories c ON c.id = r.category_id WHERE r.source = 'learned'")}
    assert patterns.get("SUPER EXAMPLE") == "Groceries"          # two lines, same category
    assert "ODD SHOP" not in patterns                            # Fun once, Gifts once: ambiguous
    assert not any("BIT" in p for p in patterns)                 # built-in transfer rule keeps priority
    from ftracker import categorize
    assert categorize.find_rule(conn, "SUPER EXAMPLE 99", -10, "card")["source"] == "learned"


def test_entries_feed_old_months_and_bank_data_takes_over_when_present(conn, sheet, home, monkeypatch):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    sheet_import.save(conn, sheet_import.parse_workbook(sheet))
    ov = summary.month_overview(conn, "2025-03")
    assert ov["income"] == 28000.0 and ov["income_source"] == "sheet"
    assert ov["spending"] >= 6741.0                              # the mortgage is counted as spending
    assert ov["put_into_savings"] == 1000.0
    run_import(home, conn, monkeypatch, [("b.pdf", bank_file(
        [(date(2025, 3, 1), "פאפאיה", 500.0)], start=date(2025, 3, 1), end=date(2025, 3, 31)))])
    from ftracker import categorize
    categorize.add_user_rule(conn, "פאפאיה", "Salary", "income", "in")
    ov = summary.month_overview(conn, "2025-03")
    assert ov["income_source"] == "bank" and ov["income"] == 500.0


def test_balances_accounts_owners_and_auto_link_to_the_real_bank_account(conn, sheet, home, monkeypatch):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    run_import(home, conn, monkeypatch, [("b.pdf", bank_file([(date(2026, 8, 1), "X", 1.0)], last4="4567"))])
    res = sheet_import.save(conn, sheet_import.parse_workbook(sheet))
    leumi = conn.execute("SELECT id FROM accounts WHERE issuer = 'leumi'").fetchone()[0]
    assert conn.execute("SELECT COUNT(*) FROM balances WHERE account_id = ?", (leumi,)).fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM accounts WHERE label = 'לאומי'").fetchone()[0] == 0   # linked, not duplicated
    owner = conn.execute("SELECT o.name FROM accounts a JOIN owners o ON o.id = a.owner_id WHERE a.label = 'שיר פנסיה כלל'").fetchone()[0]
    assert owner == "Shir"
    assert conn.execute("SELECT kind FROM accounts WHERE label = 'שיר פנסיה כלל'").fetchone()[0] == "pension"
    with pytest.raises(LookupError):
        balances.net_worth(conn, "2025-09-20")                   # dollars, no rate yet: it asks, never guesses
    fx.set_rate(conn, "2025-09-17", "USD", "ILS", 3.5, "manual")
    assert balances.net_worth(conn, "2025-09-20")["total"] > 133000 * 3.5


def test_months_in_the_history_are_not_asked_for_as_missing_statements(conn, sheet, home, monkeypatch):
    from .helpers import card_file, card_txn
    for billing in (date(2025, 2, 2), date(2025, 4, 2)):                       # March is the gap
        run_import(home, conn, monkeypatch, [(f"c{billing}.xlsx", card_file(
            billing=billing, txns=[card_txn(billing, "S", 10.0, f"v{billing}")]))], today=date(2025, 4, 10))
    before = {(i["source"], i["period"]) for i in expected.expected_files(conn, "2025-04-10")}
    assert ("Isracard 1111", "2025-03") in before
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    sheet_import.save(conn, sheet_import.parse_workbook(sheet))
    after = {(i["source"], i["period"]) for i in expected.expected_files(conn, "2025-04-10")}
    assert ("Isracard 1111", "2025-03") not in after
