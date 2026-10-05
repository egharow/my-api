from datetime import date

import openpyxl

from ftracker import accounts, balances, commits, export, fx, goals, sheet_import, summary
from .helpers import bank_file, card_file, card_txn, run_import
from .test_sheet_import import _workbook


def test_export_has_only_submitted_data_and_notes_only_on_request(conn, home, monkeypatch, tmp_path):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    run_import(home, conn, monkeypatch, [("c.xlsx", card_file(txns=[card_txn(date(2026, 9, 3), "SUPER", 100.0, "q")]))])
    out = tmp_path / "x.xlsx"
    export.export_workbook(conn, out, "2026-10-05")
    wb = openpyxl.load_workbook(out)
    assert wb["Monthly summary"].max_row == 1                       # header only: nothing submitted yet
    assert dict(wb["Info"].values)["Draft imports not included"] == 1
    commits.submit(conn, home, 1, acknowledge=True, today="2026-10-05")
    from ftracker import discrepancies as dx
    item = dx.open_manual(conn, 1, "private remark")
    conn.commit()
    export.export_workbook(conn, out, "2026-10-05")
    wb = openpyxl.load_workbook(out)
    assert [r[0] for r in wb["Monthly summary"].iter_rows(min_row=2, values_only=True)] == ["2026-10"]
    assert "Notes" not in wb.sheetnames
    export.export_workbook(conn, out, "2026-10-05", include_notes=True)
    assert "private remark" in str(list(openpyxl.load_workbook(out)["Notes"].values))


def test_export_net_worth_and_goals(conn, tmp_path):
    accounts.add_owner(conn, "Ely")
    accounts.add_account(conn, "savings", "x", "Savings", "Ely")
    accounts.add_account(conn, "investment", "x", "Brokerage", "Ely", currency="USD")
    fx.set_rate(conn, "2026-09-01", "USD", "ILS", 4.0, "manual")
    balances.set_balance(conn, "Savings", 1000, "2026-09-01")
    balances.set_balance(conn, "Brokerage", 100, "2026-09-01")
    goals.add_goal(conn, "Pot", 5000, accounts_=["Savings"])
    out = tmp_path / "x.xlsx"
    export.export_workbook(conn, out, "2026-10-05")
    wb = openpyxl.load_workbook(out)
    assert list(wb["Net worth"].iter_rows(min_row=2, values_only=True)) == [("2026-09-01", 1400.0, 350.0)]
    assert wb["Goals"]["A2"].value == "Pot" and wb["Goals"]["J2"].value in ("tracking", "no_data")


def test_summaries_can_be_limited_to_submitted_imports(conn, home, tmp_path):
    accounts.add_owner(conn, "Ely"); accounts.add_owner(conn, "Shir")
    sheet_import.save(conn, sheet_import.parse_workbook(_workbook(tmp_path / "h.xlsx")))
    assert summary.month_overview(conn, "2025-03", committed_only=True)["income"] == 0
    assert summary.month_overview(conn, "2025-03")["income"] == 28000.0
    assert summary.months_available(conn, committed_only=True) == []
    commits.submit(conn, home, 1, acknowledge=True, today="2026-10-05")
    assert summary.month_overview(conn, "2025-03", committed_only=True)["income"] == 28000.0
    months = summary.months_available(conn, committed_only=True)
    assert "2025-03" in months and months == sorted(months)
    assert {"2026-01", "2026-02"} <= set(months)          # the High level loan figures are real entries too
