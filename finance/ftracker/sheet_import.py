"""One-time import of the old budget spreadsheet (exported from Google Sheets as .xlsx).

What it reads, and where it goes:
  * each monthly tab's expenses log   -> transactions (your own categories are kept and approved)
  * each monthly tab's dashboard      -> monthly entries: income, savings, debt, budget targets
  * the "Net worth" tab               -> accounts and dated balances
  * the "High level" tab              -> monthly loan and mortgage payments

It is a preview first: nothing is written unless you save, and the original sheet is never touched.
"""
import json
import re
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

from . import categorize
from .db import audit, now
from .normalize import merchant_key, norm_description

SKIP_TABS = {"Read Me!", "Sample", "Budget", "Format", "High level", "Expenses", "Net worth"}
HEBREW_MONTHS = {"ינואר": 1, "פברואר": 2, "מרץ": 3, "אפריל": 4}

# Sheet category -> category in this app. Anything not listed is created under its own name.
CATEGORY_MAP = {
    "Food shopping": "Groceries", "Eating out": "Eating out", "Fuel/ Car": "Fuel & parking",
    "Clothes shopping": "Clothes", "Beauty care": "Beauty care", "Gifts": "Gifts", "Fun": "Fun",
    "Medical care": "Medical", "Miscellaneous": "Miscellaneous", "Memberships": "Memberships",
    "Furniture": "Household", "Household": "Household", "TV + Internet": "Internet & TV",
    "Phone": "Phone", "Electric Bill": "Utilities", "Gas": "Utilities", "Water Bill": "Utilities",
    "Education": "Education", "Insurance": "Insurance", "Arnona": "Property tax",
    "Vaad Bait": "Building committee", "Donations": "Donations",
    "Investment": "Transfer to savings",       # money moved into investments is not spending
}
LEARN_MIN_COUNT = 2
LEARN_MIN_SHARE = 0.9
LEARN_MIN_KEY = 4


@dataclass
class LogRow:
    day: date | None
    amount: float
    description: str
    category: str | None
    notes: str | None = None


@dataclass
class Entry:
    month: str
    section: str
    label: str
    expected: float | None
    actual: float | None
    source: str


@dataclass
class MonthData:
    tab: str
    month: str
    log: list[LogRow]
    entries: list[Entry]


@dataclass
class NetWorthRow:
    name: str
    section: str
    values: dict[date, tuple[float, str]]       # snapshot date -> (amount, currency)


@dataclass
class HistoryData:
    months: list[MonthData] = field(default_factory=list)
    net_worth_dates: list[date] = field(default_factory=list)
    net_worth: list[NetWorthRow] = field(default_factory=list)
    high_level: list[Entry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _num(v) -> float | None:
    if isinstance(v, bool) or v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("₪", "").replace("$", "").strip())
    except ValueError:
        return None


def _cell(row, idx):
    return row[idx] if idx < len(row) else None


def _rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _block(rows, start, label_col, exp_col, act_col, stop) -> list[tuple[str, float | None, float | None]]:
    out = []
    for r in rows[start:]:
        if stop(r):
            break
        label = _cell(r, label_col)
        if isinstance(label, str) and label.strip() and label.strip() not in ("TOTAL", "CATEGORY"):
            exp, act = _num(_cell(r, exp_col)), _num(_cell(r, act_col))
            if exp or act:
                out.append((label.strip(), exp, act))
    return out


def _parse_month_tab(title: str, ws, warnings: list[str]) -> MonthData | None:
    rows = _rows(ws)
    log_start = next((i for i, r in enumerate(rows) if _cell(r, 8) == "DATE"), None)
    if log_start is None:
        return None
    log: list[LogRow] = []
    for r in rows[log_start + 1:]:
        amount, desc = _num(_cell(r, 11)), _cell(r, 13)
        if amount is None or not isinstance(desc, str) or not desc.strip():
            continue
        d = _cell(r, 8)
        category = _cell(r, 20)
        log.append(LogRow(d.date() if isinstance(d, datetime) else None, amount, desc.strip(),
                          category.strip() if isinstance(category, str) and category.strip() else None,
                          _cell(r, 24) if isinstance(_cell(r, 24), str) else None))
    dated = [(l.day.year, l.day.month) for l in log if l.day]
    if not dated:
        return None
    y, m = Counter(dated).most_common(1)[0][0]
    month = f"{y:04d}-{m:02d}"
    undated = sum(1 for l in log if not l.day)
    if undated:
        warnings.append(f"{title}: {undated} line(s) have no date; booked on the 1st of {month}")
        for l in log:
            if not l.day:
                l.day = date(y, m, 1)

    entries: list[Entry] = []
    head = next((i for i, r in enumerate(rows) if _cell(r, 1) == "INCOME"), None)
    if head is not None:
        first = head + 2
        for label, e, a in _block(rows, first, 1, 4, 6, lambda r: _cell(r, 1) == "TOTAL"):
            entries.append(Entry(month, "income", label, e, a, title))
        sav = next((i for i, r in enumerate(rows) if _cell(r, 1) == "SAVINGS"), None)
        if sav is not None:
            for label, e, a in _block(rows, sav + 2, 1, 4, 6, lambda r: _cell(r, 1) == "TOTAL"):
                entries.append(Entry(month, "saving", label, e, a, title))
        for label, e, a in _block(rows, first, 9, 12, 14, lambda r: _cell(r, 8) == "TOTAL"):
            entries.append(Entry(month, "debt", label, e, a, title))
        for label, e, a in _block(rows, first, 17, 20, 22, lambda r: _cell(r, 16) == "TOTAL"):
            if e:
                entries.append(Entry(month, "budget", label, e, None, title))
        for label, e, a in _block(rows, first, 24, 26, 28,
                                  lambda r: _cell(r, 31) == "Bills" or _cell(r, 24) in ("TOTAL",)):
            if e:
                entries.append(Entry(month, "budget", label, e, None, title))
    return MonthData(title, month, log, entries)


def _nw_date(v, warnings: list[str]) -> date | None:
    """The sheet is day/month, but cells that look like a US date were read as month/day, so swap."""
    if isinstance(v, datetime):
        if v.day > 12:
            return v.date()
        return date(v.year, v.day, v.month)
    if isinstance(v, str):
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})", v.strip())
        if m:
            y = int(m.group(3)) + (2000 if int(m.group(3)) < 100 else 0)
            return date(y, int(m.group(2)), int(m.group(1)))
    return None


def _kind_for(name: str, section: str) -> str:
    n = name
    if section == "נדלן":
        return "real_estate"
    if section == "מניות חול":
        return "investment"
    if section == "בנקים":
        return "investment" if ("טרייד" in n or "מניות" in n) else "bank"
    if "פנסיה" in n:
        return "pension"
    return "savings"


def _parse_net_worth(ws, data: HistoryData) -> None:
    rows = _rows(ws)
    header = rows[0] if rows else ()
    cols = {}
    for idx in range(4, len(header)):
        d = _nw_date(header[idx], data.warnings)
        if d:
            cols[idx] = d
    data.net_worth_dates = sorted(cols.values())
    section = ""
    fmt = {}
    for row_cells in ws.iter_rows(min_row=2):
        r = tuple(c.value for c in row_cells)
        if _cell(r, 2) and not _cell(r, 3):
            section = str(r[2]).strip()
            continue
        name = _cell(r, 3)
        if not (isinstance(name, str) and name.strip()):
            continue
        values: dict[date, tuple[float, str]] = {}
        for idx, d in cols.items():
            cell = row_cells[idx] if idx < len(row_cells) else None
            amount = _num(cell.value) if cell is not None else None
            if amount is None:
                continue
            number_format = cell.number_format or ""
            currency = "USD" if ("$" in number_format and "₪" not in number_format) else "ILS"
            if number_format == "General":
                data.warnings.append(f"Net worth: {name.strip()} on {d} has no currency format; assumed ₪{amount:,.0f}")
            values[d] = (amount, currency)
        if values:
            data.net_worth.append(NetWorthRow(name.strip(), section, values))


def _parse_high_level(ws, year: int, data: HistoryData) -> None:
    rows = _rows(ws)
    header = next((r for r in rows if any(c in HEBREW_MONTHS for c in r if isinstance(c, str))), None)
    if not header:
        return
    months = {i: HEBREW_MONTHS[c] for i, c in enumerate(header) if isinstance(c, str) and c in HEBREW_MONTHS}
    section, n = None, 0
    for r in rows[rows.index(header) + 1:]:
        label = r[0] if r and isinstance(r[0], str) and r[0].strip() else None
        if label:
            section, n = label.strip(), 0
        if section not in ("הלוואות", "משכנתא"):
            continue
        n += 1
        name = "Mortgage" if section == "משכנתא" else f"Loan {n}"
        for idx, mo in months.items():
            v = _num(_cell(r, idx))
            if v:
                data.high_level.append(Entry(f"{year:04d}-{mo:02d}", "debt", name, None, v, "High level"))
    if data.high_level:
        data.warnings.append(f"High level tab has no year; assumed {year} (its loan amounts match your 2026 bank statement)")


def parse_workbook(path, high_level_year: int = 2026) -> HistoryData:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    data = HistoryData()
    seen: dict[str, str] = {}
    for ws in wb:
        if ws.title in SKIP_TABS:
            continue
        month = _parse_month_tab(ws.title, ws, data.warnings)
        if not month:
            continue
        if month.month in seen:
            data.warnings.append(f"{ws.title} and {seen[month.month]} are both {month.month}; kept the first, skipped this one")
            continue
        seen[month.month] = ws.title
        data.months.append(month)
    data.months.sort(key=lambda m: m.month)
    if "Net worth" in wb.sheetnames:
        _parse_net_worth(wb["Net worth"], data)
    if "High level" in wb.sheetnames:
        _parse_high_level(wb["High level"], high_level_year, data)
    if "Expenses" in wb.sheetnames:
        data.warnings.append("The 'Expenses' tab holds a pasted card statement whose billing date and purchase dates "
                             "disagree; it was not imported. Upload the original statement instead.")
    return data


def unmapped_categories(data: HistoryData, extra_map: dict[str, str] | None = None) -> Counter:
    mapping = {**CATEGORY_MAP, **(extra_map or {})}
    return Counter(l.category for m in data.months for l in m.log if l.category and l.category not in mapping)


def preview(data: HistoryData, extra_map: dict[str, str] | None = None) -> str:
    out = []
    n_tx = sum(len(m.log) for m in data.months)
    out.append(f"Monthly tabs: {len(data.months)} ({data.months[0].month} to {data.months[-1].month}), {n_tx} log lines"
               if data.months else "No monthly tabs found")
    for m in data.months:
        total = sum(l.amount for l in m.log)
        blank = sum(1 for l in m.log if not l.category)
        out.append(f"  {m.month}  tab “{m.tab}”: {len(m.log)} lines, ₪{total:,.2f}, {blank} without a category, "
                   f"{len([e for e in m.entries if e.section == 'income'])} income source(s)")
    gaps = _missing_months(data)
    if gaps:
        out.append(f"No tab for: {', '.join(gaps)}")
    un = unmapped_categories(data, extra_map)
    if un:
        out.append("Categories with no mapping (they will be created under the same name; use --map to change):")
        for name, n in un.most_common():
            out.append(f"  {name!r}: {n} line(s)")
    out.append(f"Net worth: {len(data.net_worth)} accounts, snapshots on "
               + ", ".join(d.isoformat() for d in data.net_worth_dates))
    out.append("  (dates were read as day/month: the sheet's US-looking dates were swapped back)")
    if data.high_level:
        out.append(f"High level: {len(data.high_level)} loan/mortgage figures")
    for w in data.warnings:
        out.append(f"! {w}")
    return "\n".join(out)


def _missing_months(data: HistoryData) -> list[str]:
    have = {m.month for m in data.months}
    if not have:
        return []
    y, m = map(int, min(have).split("-"))
    ey, em = map(int, max(have).split("-"))
    out = []
    while (y, m) <= (ey, em):
        key = f"{y:04d}-{m:02d}"
        if key not in have:
            out.append(key)
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def _category_id(conn, name: str) -> int:
    row = conn.execute("SELECT id FROM categories WHERE name = ?", (name,)).fetchone()
    if row:
        return row[0]
    return conn.execute("INSERT INTO categories (name, kind, neutral) VALUES (?, 'expense', 0)", (name,)).lastrowid


def _owner(conn, name: str | None):
    if not name:
        return None
    row = conn.execute("SELECT id FROM owners WHERE name = ? COLLATE NOCASE", (name,)).fetchone()
    return row[0] if row else None


def _history_account(conn) -> int:
    row = conn.execute("SELECT id FROM accounts WHERE issuer = 'sheet' AND last4 = 'hist'").fetchone()
    if row:
        return row[0]
    return conn.execute(
        "INSERT INTO accounts (kind, issuer, label, last4, currency, created_at) "
        "VALUES ('other','sheet','Sheet history (all cards)','hist','ILS',?)", (now(),)).lastrowid


def save(conn: sqlite3.Connection, data: HistoryData, primary_owner: str = "Ely", partner_owner: str = "Shir",
         extra_map: dict[str, str] | None = None, links: dict[str, str] | None = None) -> dict:
    """Write everything as one draft import. Safe to run twice: existing rows are skipped."""
    mapping = {**CATEGORY_MAP, **(extra_map or {})}
    links = links or {}
    batch_id = conn.execute("INSERT INTO batches (created_at, note) VALUES (?, 'history import from the old budget sheet')",
                            (now(),)).lastrowid
    acct = _history_account(conn)
    new_tx = skipped = overridden = 0
    for m in data.months:
        seen: Counter = Counter()
        for l in m.log:
            norm = norm_description(l.description)
            sig = (l.day, round(l.amount, 2), norm)
            key = f"sheet|{m.month}|{l.day}|{l.amount:.2f}|{norm}|{seen[sig]}"
            seen[sig] += 1
            cat_name = mapping.get(l.category, l.category) if l.category else None
            special = categorize.find_rule(conn, norm, -l.amount, "card")
            if (cat_name is None and special and special["source"] == "builtin"
                    and special["set_kind"] in ("transfer", "card_payment")):
                overridden += 1      # only for lines you left blank: Bit usually moves money, it is not spending
            try:
                cur = conn.execute(
                    """INSERT INTO transactions (batch_id, account_id, dedupe_key, txn_date, budget_month, description,
                              description_norm, detail, amount, currency, category_id, category_status, confidence,
                              proposal_basis, kind, created_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (batch_id, acct, key, l.day.isoformat(), m.month, l.description, norm, l.notes, -l.amount, "ILS",
                     _category_id(conn, cat_name) if cat_name else categorize.uncategorised_id(conn),
                     "approved" if cat_name else "proposed", 1.0 if cat_name else 0.0,
                     "from your old sheet" if cat_name else "no category in your old sheet",
                     "refund" if l.amount < 0 else "purchase", now()))
            except sqlite3.IntegrityError:
                skipped += 1
                continue
            new_tx += 1
            if not cat_name:
                categorize.classify(conn, cur.lastrowid)
        for e in m.entries:
            _save_entry(conn, e, batch_id, primary_owner, partner_owner)
    for e in data.high_level:
        _save_entry(conn, e, batch_id, primary_owner, partner_owner)

    conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('history_months', ?)",
                 (json.dumps(sorted({m.month for m in data.months})),))
    nw = _save_net_worth(conn, data, batch_id, primary_owner, partner_owner, links)
    learned = learn_rules(conn)
    audit(conn, "sheet_import", "batch", batch_id, f"{new_tx} lines, {nw['balances']} balances", "import")
    if new_tx == 0 and not conn.execute("SELECT 1 FROM monthly_entries WHERE batch_id = ?", (batch_id,)).fetchone():
        conn.execute("DELETE FROM balances WHERE batch_id = ?", (batch_id,))
        conn.execute("DELETE FROM batches WHERE id = ?", (batch_id,))   # nothing new: leave no empty draft behind
        batch_id = None
    conn.commit()
    return {"batch_id": batch_id, "transactions": new_tx, "already_there": skipped, "rules_learned": learned, "transfers_reclassified": overridden, **nw}


def _save_entry(conn, e: Entry, batch_id: int, primary: str, partner: str) -> None:
    owner = _owner(conn, e.label if e.label in (primary, partner) else None)
    conn.execute(
        """INSERT INTO monthly_entries (month, section, label, owner_id, expected, actual, source, batch_id, created_at)
           VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(month, section, label, source)
           DO UPDATE SET expected = excluded.expected, actual = excluded.actual""",
        (e.month, e.section, e.label, owner, e.expected, e.actual, e.source, batch_id, now()))


def _save_net_worth(conn, data: HistoryData, batch_id: int, primary: str, partner: str, links: dict[str, str]) -> dict:
    created = balances = 0
    for row in data.net_worth:
        owner = _owner(conn, partner if row.section.startswith("שיר") else primary)
        ref = links.get(row.name)
        if ref is None and row.name == "לאומי":
            hit = conn.execute("SELECT id FROM accounts WHERE issuer = 'leumi' AND kind = 'bank'").fetchall()
            ref = hit[0][0] if len(hit) == 1 else None
        if isinstance(ref, str):
            from .accounts import find_account
            ref = find_account(conn, ref)["id"]
        if ref is None:
            existing = conn.execute("SELECT id FROM accounts WHERE issuer = 'manual' AND label = ?", (row.name,)).fetchone()
            if existing:
                ref = existing[0]
            else:
                currency = next(iter(row.values.values()))[1]
                ref = conn.execute(
                    "INSERT INTO accounts (kind, issuer, label, owner_id, currency, created_at) VALUES (?,?,?,?,?,?)",
                    (_kind_for(row.name, row.section), "manual", row.name, owner, currency, now())).lastrowid
                created += 1
        for d, (amount, currency) in row.values.items():
            conn.execute(
                """INSERT INTO balances (account_id, as_of, amount, currency, batch_id, note, created_at)
                   VALUES (?,?,?,?,?,?,?) ON CONFLICT(account_id, as_of) DO UPDATE SET amount = excluded.amount""",
                (ref, d.isoformat(), amount, currency, batch_id, "from your old sheet", now()))
            balances += 1
    return {"accounts_created": created, "balances": balances}


def learn_rules(conn: sqlite3.Connection) -> int:
    """Turn consistent categorisations in your history into rules for new statements.

    A merchant becomes a rule only if it appears at least twice and one category covers 90%+ of
    its lines. Merchants that built-in rules treat specially (Bit transfers, card payments, fees)
    are skipped so history cannot override them.
    """
    groups: dict[str, Counter] = defaultdict(Counter)
    for r in conn.execute(
            """SELECT t.description, c.name FROM transactions t JOIN categories c ON c.id = t.category_id
               WHERE t.category_status = 'approved' AND t.proposal_basis = 'from your old sheet'"""):
        key = merchant_key(r["description"])
        if len(key) >= LEARN_MIN_KEY:
            groups[key][r["name"]] += 1
    created = 0
    for key, counts in groups.items():
        total = sum(counts.values())
        name, top = counts.most_common(1)[0]
        if total < LEARN_MIN_COUNT or top / total < LEARN_MIN_SHARE:
            continue
        builtin = categorize.find_rule(conn, key, -1.0, "card")
        if builtin and builtin["source"] == "builtin" and builtin["set_kind"]:
            continue
        before = conn.execute("SELECT COUNT(*) FROM rules WHERE pattern = ? AND source = 'learned'", (key,)).fetchone()[0]
        categorize.add_learned_rule(conn, key, name)
        created += 0 if before else 1
    return created


def restore_categories(conn: sqlite3.Connection, data: HistoryData, extra_map: dict[str, str] | None = None) -> int:
    """Put your own categories back on imported lines that an earlier version changed.

    Only lines in imports that are not submitted yet, and not ones you have since edited, are touched.
    """
    mapping = {**CATEGORY_MAP, **(extra_map or {})}
    fixed = 0
    for m in data.months:
        seen: Counter = Counter()
        for l in m.log:
            norm = norm_description(l.description)
            sig = (l.day, round(l.amount, 2), norm)
            key = f"sheet|{m.month}|{l.day}|{l.amount:.2f}|{norm}|{seen[sig]}"
            seen[sig] += 1
            if not l.category:
                continue
            want = mapping.get(l.category, l.category)
            row = conn.execute(
                """SELECT t.id, c.name AS cat, t.proposal_basis, b.status FROM transactions t
                   JOIN categories c ON c.id = t.category_id JOIN batches b ON b.id = t.batch_id
                   WHERE t.dedupe_key = ?""", (key,)).fetchone()
            if not row or row["cat"] == want or row["status"] == "committed" or row["proposal_basis"] == "approved by you":
                continue
            conn.execute(
                """UPDATE transactions SET category_id = ?, category_status = 'approved', confidence = 1.0,
                          proposal_basis = 'from your old sheet', kind = ? WHERE id = ?""",
                (_category_id(conn, want), "refund" if l.amount < 0 else "purchase", row["id"]))
            fixed += 1
    return fixed
