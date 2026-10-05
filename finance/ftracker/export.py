"""A clean workbook for the shared Google Sheet: committed numbers only, never drafts."""
import sqlite3
from datetime import datetime
from pathlib import Path

from . import balances, discrepancies, fx, goals, summary


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    from openpyxl.styles import Font
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(
            max((len(str(c.value)) for c in col if c.value is not None), default=8) + 2, 48)
    return ws


def export_workbook(conn: sqlite3.Connection, path: Path, today: str, include_notes: bool = False) -> dict:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    months = summary.months_available(conn, committed_only=True)

    last = conn.execute("SELECT committed_at FROM commits ORDER BY committed_at DESC LIMIT 1").fetchone()
    drafts = conn.execute("SELECT COUNT(*) FROM batches WHERE status != 'committed'").fetchone()[0]
    _sheet(wb, "Info", ["Item", "Value"], [
        ("Generated", datetime.now().isoformat(timespec="seconds")),
        ("Latest submitted import", last[0] if last else "none yet"),
        ("Draft imports not included", drafts),
        ("Currency", "ILS unless a column says otherwise"),
    ])

    rows = []
    for m in months:
        ov = summary.month_overview(conn, m, committed_only=True)
        rows.append((m, ov["income"], ov["spending"], ov["saved"], ov["put_into_savings"], ov["income_source"]))
    _sheet(wb, "Monthly summary", ["Month", "Income", "Spending", "Left over", "Put into savings", "Income from"], rows)

    owners = [None] + [r[0] for r in conn.execute("SELECT name FROM owners ORDER BY name")]
    cat_rows = []
    for m in months:
        for owner in owners:
            for r in summary.spending_by_category(conn, m, owner, committed_only=True):
                cat_rows.append((m, owner or "Household", r["category"], r["subcategory"], r["spent"], r["n"]))
    _sheet(wb, "By category", ["Month", "Who", "Category", "Subcategory", "Spent", "Lines"], cat_rows)

    nw_rows = []
    for d in [r[0] for r in conn.execute("SELECT DISTINCT as_of FROM balances ORDER BY as_of")]:
        try:
            ils = balances.net_worth(conn, d, "ILS")["total"]
            usd = balances.net_worth(conn, d, "USD")["total"]
        except LookupError:
            ils = usd = None
        nw_rows.append((d, ils, usd))
    _sheet(wb, "Net worth", ["Date", "Net worth ILS", "Net worth USD"], nw_rows)

    acc_rows = [(r["label"], r["kind"], r["owner"], r["as_of"], r["amount"], r["currency"]) for r in balances.latest(conn, today)]
    _sheet(wb, "Accounts", ["Account", "Type", "Owner", "As of", "Balance", "Currency"], acc_rows)

    goal_rows = []
    for p in goals.all_progress(conn, today):
        goal_rows.append((p["name"], p["target"], p["currency"], p["current"], p["percent"], p["target_date"], p["required_per_month"],
                          p["growth_per_month"], p["projected_date"], p["status"]))
    _sheet(wb, "Goals", ["Goal", "Target", "Currency", "Now", "Progress", "Target date", "Needed per month",
                         "Growing per month", "Reached about", "Status"], goal_rows)

    if include_notes:
        notes = []
        for d in discrepancies.list_items(conn, ("open", "explained", "acknowledged")):
            thread = "; ".join(f"{c['author']}: {c['body']}" for c in discrepancies.thread(conn, d["id"]))
            notes.append((d["status"], d["title"], thread, d["follow_up_date"]))
        _sheet(wb, "Notes", ["Status", "Item", "Comments", "Follow up"], notes)
    wb.save(path)
    return {"months": len(months), "included_notes": include_notes, "drafts_excluded": drafts}
