import re
from datetime import date

from .base import (ParseError, ParsedFile, ParsedStatement, ParsedTxn, apply_detail, clean,
                   currency_code, parse_date, to_float)

HEBREW_MONTHS = {
    "ינואר": 1, "פברואר": 2, "מרץ": 3, "אפריל": 4, "מאי": 5, "יוני": 6,
    "יולי": 7, "אוגוסט": 8, "ספטמבר": 9, "אוקטובר": 10, "נובמבר": 11, "דצמבר": 12,
}
_DATE = re.compile(r"\d{1,2}\.\d{1,2}\.\d{2,4}$")


def read_rows(path) -> list[list]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def looks_like(rows: list[list]) -> bool:
    head = " ".join(clean(c) for r in rows[:8] for c in r if c)
    return "פירוט עסקאות" in head and ("מסטרקארד" in head or "ויזה" in head or "ישראכרט" in head)


def parse_rows(rows: list[list]) -> ParsedFile:
    card_line = month_year = holder = billing_text = None
    header_idx = None
    for i, r in enumerate(rows):
        cells = [clean(c) for c in r]
        text = " ".join(c for c in cells if c)
        if card_line is None and re.search(r"-\s*\d{4}\s*$", cells[0] if cells else ""):
            card_line = cells[0]
        if month_year is None:
            for c in cells:
                m = re.fullmatch(r"(\S+)\s+(\d{4})", c)
                if m and m.group(1) in HEBREW_MONTHS:
                    month_year = (HEBREW_MONTHS[m.group(1)], int(m.group(2)))
        if cells and cells[0].startswith("על שם"):
            holder = cells[0][len("על שם"):].strip()
            for c in cells[1:]:
                if "לחיוב ב" in c:
                    billing_text = c
        if "לחיוב ב" in text and billing_text is None:
            billing_text = next(c for c in cells if "לחיוב ב" in c)
        if cells and cells[0] == "תאריך רכישה" and header_idx is None:
            header_idx = i
    if not (card_line and header_idx is not None):
        raise ParseError("not an Isracard statement")

    last4 = re.search(r"(\d{4})\s*$", card_line).group(1)
    product = card_line.rsplit("-", 1)[0].strip()

    billing = None
    if billing_text and month_year:
        m = re.search(r"(\d{1,2})\.(\d{1,2})", billing_text)
        if m:
            day, month = int(m.group(1)), int(m.group(2))
            # The month on the sheet is the statement month; the debit can land in it.
            year = month_year[1] + (1 if month < month_year[0] else 0)
            billing = date(year, month, day)
    if billing is None:
        raise ParseError("could not read the billing date")

    txns: list[ParsedTxn] = []
    stated_total = None
    for r in rows[header_idx + 1:]:
        cells = list(r) + [None] * 8
        first = clean(cells[0])
        merchant = clean(cells[1])
        if _DATE.match(first):
            charge_cur = currency_code(cells[5])
            orig_cur = currency_code(cells[3])
            orig = to_float(cells[2])
            charge = to_float(cells[4])
            t = ParsedTxn(
                txn_date=parse_date(first), description=merchant,
                amount=-charge, currency=charge_cur,
                orig_amount=-orig, orig_currency=orig_cur,
                voucher=clean(cells[6]) or None,
            )
            apply_detail(t, cells[7])
            txns.append(t)
        elif merchant.startswith("סה\"כ לחיוב"):
            stated_total = to_float(cells[4])
    if stated_total is None:
        raise ParseError("statement total row not found")

    st = ParsedStatement(
        issuer="isracard", account_kind="card", last4=last4, product=product,
        billing_date=billing, period_start=None, period_end=None,
        stated_total=stated_total, currency="ILS", txns=txns, holder_text=holder or "",
    )
    return ParsedFile("isracard", "card_export", [st])


def parse(path) -> ParsedFile:
    return parse_rows(read_rows(path))
